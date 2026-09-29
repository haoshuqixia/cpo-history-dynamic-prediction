from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path

import numpy as np
import pandas as pd

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ModuleNotFoundError:
    Document = None


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = Path(__file__).resolve().parent
FIGDIR = OUTDIR / "figures_r1"
OUTDOC = OUTDIR / "Supplementary_Information_CPO_history_V1_1_R1.docx"

BLUE = "#2F89C5"
GREY = "#7A7A7A"
LIGHT_GREY = "#D8D8D8"
PURPLE = "#7856A8"
BLACK = "#111111"


def p(*parts: str) -> Path:
    return ROOT.joinpath(*parts)


MIMIC_WEIGHT_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_STAGE1A_FEATURE_AND_CENSORING_20260908",
    "results",
    "observation_weights_v1.csv.gz",
)
EICU_WEIGHT_FILE = p(
    "CPO_EICU_EXTERNAL_DISCHARGE_BOUNDARY_R1_V1_0",
    "results",
    "stage_e1b_observation_weights_final.csv.gz",
)
PREDICTOR_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_PAPER_PACKAGE_20260908",
    "tables",
    "tableS3_predictor_dictionary_and_missingness.csv",
)
MAIN_METRICS_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_STAGE1B_ABC_TEMPORAL_20260908",
    "results",
    "temporal_metrics.csv",
)
MAIN_DELTAS_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_STAGE1B_ABC_TEMPORAL_20260908",
    "results",
    "paired_temporal_deltas.csv",
)
R1_DELTAS_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_REVIEWER_DEFENSE_20260908",
    "stage_r1_storetime",
    "results",
    "paired_temporal_deltas.csv",
)
R2_DELTAS_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_REVIEWER_DEFENSE_20260908",
    "stage_r2_robustness",
    "results",
    "paired_temporal_deltas.csv",
)
CLEAN_METRICS_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_STAGE1C_EARLY_WARNING_AUDIT_20260908",
    "results",
    "clean_history_temporal_metrics.csv",
)
SUBSET_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_STAGE1C_EARLY_WARNING_AUDIT_20260908",
    "results",
    "locked_prediction_interpretation_subsets.csv",
)
EICU_ABS_FILE = p(
    "CPO_EICU_EXTERNAL_DISCHARGE_BOUNDARY_R1_V1_0",
    "results",
    "stage_e2_absolute_metrics.csv",
)
HOSPITAL_FILE = p(
    "CPO_EICU_EXTERNAL_DISCHARGE_BOUNDARY_R1_V1_0",
    "results",
    "stage_e2_hospital_heterogeneity.csv",
)
ALERT_FILE = p(
    "CPO_STATISTICS_RERUN_20260915",
    "CV_PAC_CPO_STAGE1D_ALERT_UTILITY_20260908",
    "results",
    "alert_utility.csv",
)


def set_mpl():
    import matplotlib as mpl
    import matplotlib.pyplot as plt

    mpl.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans", "sans-serif"],
            "font.size": 7,
            "axes.labelsize": 7,
            "axes.titlesize": 8,
            "xtick.labelsize": 6.5,
            "ytick.labelsize": 6.5,
            "legend.fontsize": 6.5,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.7,
            "pdf.fonttype": 42,
            "svg.fonttype": "none",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )
    return plt


def save_figure(fig, stem: str) -> Path:
    import matplotlib.pyplot as plt

    FIGDIR.mkdir(parents=True, exist_ok=True)
    png = FIGDIR / f"{stem}.png"
    fig.savefig(png, dpi=600, bbox_inches="tight", facecolor="white")
    fig.savefig(FIGDIR / f"{stem}.pdf", bbox_inches="tight", facecolor="white")
    fig.savefig(FIGDIR / f"{stem}.svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return png


def prepare_weight_data() -> dict[str, pd.DataFrame]:
    mimic = pd.read_csv(MIMIC_WEIGHT_FILE)
    mimic = mimic[mimic["primary_eligible_at_landmark"].astype(bool)].copy()
    development = mimic[mimic["interval_conservative_split"].str.contains("development")].copy()
    temporal = mimic[mimic["interval_conservative_split"].str.contains("temporal")].copy()
    eicu = pd.read_csv(EICU_WEIGHT_FILE)
    return {"Development": development, "Temporal": temporal, "eICU": eicu}


def fig_s1(weight_data: dict[str, pd.DataFrame]) -> Path:
    plt = set_mpl()
    fig, axes = plt.subplots(2, 3, figsize=(7.15, 4.25), constrained_layout=True)
    names = ["Development", "Temporal", "eICU"]
    letters = list("abcdef")
    for j, name in enumerate(names):
        d = weight_data[name]
        ax = axes[0, j]
        probs = d["p_outcome_observed_6h"].dropna().to_numpy()
        ax.hist(probs, bins=np.linspace(0, 1, 31), density=True, color=BLUE, alpha=0.85, edgecolor="white", linewidth=0.25)
        ax.set_xlim(0, 1)
        ax.set_xlabel("Predicted probability of outcome observation")
        if j == 0:
            ax.set_ylabel("Density")
        ax.set_title(name, pad=4, weight="bold")
        ax.text(-0.16, 1.06, letters[j], transform=ax.transAxes, fontsize=8.5, fontweight="bold", va="top")
        ax.grid(axis="y", color="#ECECEC", linewidth=0.5)

        axw = axes[1, j]
        obs_col = "primary_outcome_observed_6h" if name != "eICU" else "outcome_observed_6h"
        observed = d[d[obs_col].astype(bool)].copy()
        raw = observed["ipow_stabilized"].dropna().to_numpy()
        trunc = observed["ipow_stabilized_truncated"].dropna().to_numpy()
        q01, q99 = np.quantile(raw, [0.01, 0.99])
        lo = max(0.1, min(raw.min(), trunc.min()) * 0.9)
        hi = max(raw.max(), trunc.max()) * 1.05
        bins = np.geomspace(lo, hi, 42)
        axw.hist(raw, bins=bins, density=True, histtype="stepfilled", color=LIGHT_GREY, alpha=0.75, label="Raw")
        axw.hist(trunc, bins=bins, density=True, histtype="step", color=BLUE, linewidth=1.4, label="Truncated")
        axw.axvline(q01, color="#555555", linestyle="--", linewidth=0.8)
        axw.axvline(q99, color="#555555", linestyle="--", linewidth=0.8)
        axw.set_xscale("log")
        from matplotlib.ticker import FixedLocator, FuncFormatter
        if hi <= 3:
            ticks = [0.7, 1.0, 2.0]
        elif hi <= 8:
            ticks = [0.6, 1.0, 2.0, 5.0]
        else:
            ticks = [0.6, 1.0, 2.0, 5.0, 10.0]
        ticks = [t for t in ticks if lo <= t <= hi]
        axw.xaxis.set_major_locator(FixedLocator(ticks))
        axw.xaxis.set_major_formatter(FuncFormatter(lambda x, _pos: f"{x:g}"))
        axw.xaxis.set_minor_formatter(FuncFormatter(lambda _x, _pos: ""))
        axw.set_xlabel("Stabilised observation weight (log scale)")
        if j == 0:
            axw.set_ylabel("Density")
        ess = (trunc.sum() ** 2) / np.square(trunc).sum()
        axw.text(0.98, 0.94, f"ESS = {ess:,.0f} ({100 * ess / len(trunc):.1f}%)", transform=axw.transAxes, ha="right", va="top", fontsize=6.6)
        axw.text(-0.16, 1.06, letters[j + 3], transform=axw.transAxes, fontsize=8.5, fontweight="bold", va="top")
        axw.grid(axis="y", color="#ECECEC", linewidth=0.5)
        if j == 2:
            axw.legend(loc="upper left")
    return save_figure(fig, "figure_S1_observation_weight_diagnostics")


def read_delta(path: Path, comparison: str, metric: str, analysis: str | None = None) -> tuple[float, float, float]:
    d = pd.read_csv(path)
    q = d[(d["comparison"] == comparison) & (d["metric"] == metric)]
    if analysis is not None:
        q = q[q["analysis"] == analysis]
    r = q.iloc[0]
    return float(r["estimate"]), float(r["ci_low"]), float(r["ci_high"])


def sensitivity_rows() -> list[dict[str, object]]:
    clean = pd.read_csv(CLEAN_METRICS_FILE)
    subset = pd.read_csv(SUBSET_FILE)

    def clean_delta(metric: str) -> tuple[float, float, float]:
        r = clean[(clean["model"] == "B-A") & (clean["metric"] == metric)].iloc[0]
        return float(r.estimate), float(r.ci_low), float(r.ci_high)

    s70 = subset[subset["subset"] == "current_cpo_ge_0.70"].set_index("model")
    s70_delta = {
        "brier": float(s70.loc["B", "brier"] - s70.loc["A", "brier"]),
        "auroc": float(s70.loc["B", "auroc"] - s70.loc["A", "auroc"]),
        "average_precision": float(s70.loc["B", "average_precision"] - s70.loc["A", "average_precision"]),
        "log_loss": float(s70.loc["B", "log_loss"] - s70.loc["A", "log_loss"]),
    }

    rows: list[dict[str, object]] = []
    specs = [
        ("Primary temporal analysis", 3416, 202, 62, MAIN_DELTAS_FILE, None),
        ("No prior CPO <0.60 W", 3117, 137, 46, CLEAN_METRICS_FILE, "clean"),
        ("Latest CPO ≥0.70 W", 3089, 131, 56, None, "point"),
        ("Strict information clock", 2266, 117, 46, R1_DELTAS_FILE, None),
        ("Broader CPO-eligible risk set", 3426, 208, 66, R2_DELTAS_FILE, "expanded_cpo_only_weighted"),
        ("Unweighted complete observation", 3416, 202, 62, R2_DELTAS_FILE, "locked_common_unweighted"),
    ]
    for label, landmarks, events, patients, path, mode in specs:
        row: dict[str, object] = {"label": label, "landmarks": landmarks, "events": events, "patients": patients}
        for metric in ["brier", "auroc", "average_precision", "log_loss"]:
            if mode == "clean":
                row[metric] = clean_delta(metric)
            elif mode == "point":
                row[metric] = (s70_delta[metric], math.nan, math.nan)
            else:
                row[metric] = read_delta(path, "B-A", metric, mode)
        rows.append(row)
    return rows


def fig_s2(rows: list[dict[str, object]]) -> Path:
    plt = set_mpl()
    fig, axes = plt.subplots(1, 2, figsize=(7.15, 3.25), sharey=True, constrained_layout=True)
    y = np.arange(len(rows))[::-1]
    labels = [str(r["label"]) for r in rows]
    configs = [
        ("brier", "Delta Brier score (Model B - Model A)", "Left favours Model B", "a"),
        ("auroc", "Delta AUROC (Model B - Model A)", "Right favours Model B", "b"),
    ]
    for ax, (metric, xlabel, direction, letter) in zip(axes, configs):
        vals = np.array([r[metric][0] for r in rows], dtype=float)
        lows = np.array([r[metric][1] for r in rows], dtype=float)
        highs = np.array([r[metric][2] for r in rows], dtype=float)
        for i, yy in enumerate(y):
            primary = i == 0
            if np.isfinite(lows[i]):
                ax.errorbar(
                    vals[i], yy,
                    xerr=[[vals[i] - lows[i]], [highs[i] - vals[i]]],
                    fmt="o", color=BLUE, ecolor=BLUE,
                    markersize=5.5 if primary else 4.2,
                    linewidth=1.15, capsize=2.2,
                )
            else:
                ax.plot(vals[i], yy, marker="o", markersize=5, markerfacecolor="white", markeredgecolor=BLUE, markeredgewidth=1.2)
        ax.axvline(0, color=BLACK, linewidth=0.9)
        ax.set_yticks(y, labels)
        ax.set_xlabel(xlabel)
        ax.set_title(direction, loc="right", fontsize=6.5, color="#444444", pad=5)
        ax.text(-0.12, 1.08, letter, transform=ax.transAxes, fontsize=8.5, fontweight="bold", va="top")
        ax.grid(axis="x", color="#ECECEC", linewidth=0.5)
        from matplotlib.ticker import MaxNLocator
        ax.xaxis.set_major_locator(MaxNLocator(nbins=5))
    axes[0].get_yticklabels()[0].set_fontweight("bold")
    axes[1].tick_params(labelleft=False)
    fig.text(0.5, -0.02, "Open markers indicate point estimates for which a frozen bootstrap interval was not available.", ha="center", fontsize=6.5)
    return save_figure(fig, "figure_S2_robustness_forest")


def fig_s3() -> Path:
    plt = set_mpl()
    metrics = [
        ("brier", "Brier score"),
        ("auroc", "AUROC"),
        ("average_precision", "Average precision"),
        ("log_loss", "Log loss"),
    ]
    fig, axes = plt.subplots(1, 4, figsize=(7.15, 1.95), constrained_layout=True)
    for ax, (metric, title) in zip(axes, metrics):
        est, low, high = read_delta(MAIN_DELTAS_FILE, "C-B", metric)
        pad = max((high - low) * 0.35, abs(est) * 0.35, 0.0002)
        ax.errorbar(est, 0, xerr=[[est - low], [high - est]], fmt="o", color=PURPLE, ecolor=PURPLE, markersize=5, linewidth=1.2, capsize=2.5)
        ax.axvline(0, color=BLACK, linewidth=0.9)
        ax.set_xlim(low - pad, high + pad)
        ax.set_ylim(-0.8, 0.8)
        ax.set_yticks([])
        ax.set_title(title, fontweight="bold", pad=7)
        ax.set_xlabel("Model C - Model B")
        ax.grid(axis="x", color="#ECECEC", linewidth=0.5)
        from matplotlib.ticker import MaxNLocator
        ax.xaxis.set_major_locator(MaxNLocator(nbins=4))
        ax.text(0.5, 0.82, f"{est:+.4f}\n({low:+.4f} to {high:+.4f})", transform=ax.transAxes, ha="center", va="top", fontsize=6.3)
    return save_figure(fig, "figure_S3_model_c_incremental_value")


def set_font(run, name: str = "Arial", size: float | None = None, bold: bool | None = None, italic: bool | None = None, color: str = BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edges[edge].items():
            element.set(qn("w:" + key), str(value))


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_font(run, size=8, color="555555")


def add_landscape_section(doc: Document) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.58)
    sec.left_margin = Inches(0.58)
    sec.right_margin = Inches(0.58)
    sec.footer.is_linked_to_previous = True


def add_caption(doc: Document, number: str, title: str) -> None:
    para = doc.add_paragraph()
    para.style = doc.styles["Heading 1"]
    para.paragraph_format.space_after = Pt(7)
    r = para.add_run(f"Supplementary {number} | {title}")
    set_font(r, size=11, bold=True)


def add_note(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.05
    r = para.add_run(text)
    set_font(r, size=8, color="333333")


def add_nature_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], numeric_cols: set[int], bold_first_rows: set[int] | None = None, font_size: float = 7.4) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)

    head = table.rows[0]
    set_repeat_table_header(head)
    for j, (cell, text) in enumerate(zip(head.cells, headers)):
        cell.width = Inches(widths[j])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=85, bottom=85, start=60, end=60)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        r = para.add_run(text)
        set_font(r, size=font_size, bold=True)
        set_cell_border(cell, top={"val": "single", "sz": "10", "color": "000000"}, bottom={"val": "single", "sz": "8", "color": "000000"})

    bold_first_rows = bold_first_rows or set()
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=65, bottom=65, start=60, end=60)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            r = para.add_run(str(value))
            set_font(r, size=font_size, bold=(i in bold_first_rows and j == 0))
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "10", "color": "000000"})


def fmt_delta(t: tuple[float, float, float], decimals: int) -> str:
    est, low, high = t
    if not (np.isfinite(low) and np.isfinite(high)):
        return f"{est:+.{decimals}f} (—)"
    return f"{est:+.{decimals}f} ({low:+.{decimals}f} to {high:+.{decimals}f})"


def make_table_s1() -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    headers = ["Variable", "MIMIC-IV source", "eICU source", "Derivation / mapping", "Time rule", "Notes"]
    rows = [
        ["Cardiac output", "chartevents item 224842 (CCO)", "nurseCharting exact label/name CO|CO", "Direct measurement", "Measurement charttime in MIMIC; entry offset available by landmark in eICU", "eICU does not distinguish continuous thermodilution, intermittent thermodilution, or Fick CO"],
        ["MAP", "chartevents items 220052, 225312, and 220181", "Exact-offset invasive nurse MAP; otherwise preceding <=5 min vitalPeriodic.systemicMean; otherwise exact-offset vitalAperiodic.noninvasiveMean", "Prespecified source-priority rule", "Concurrent in MIMIC; no future pairing in eICU", "Valid range 20-250 mmHg"],
        ["CPO", "Derived", "Derived", "CO x MAP / 451", "At paired CO-MAP time", "Native CPO was not the primary source"],
        ["mPAP", "chartevents item 220061", "nurseCharting exact label/name PA|PA Mean", "Direct mapped measurement", "Current value and summaries in (L-4 h, L]", "Used dynamically only by Model C"],
        ["SvO₂", "chartevents item 223772", "nurseCharting exact label/name SVO2|SVO2", "Direct mapped measurement", "Current value and summaries in (L-4 h, L]", "MIMIC ScvO₂ item 226541 was not substituted"],
    ]
    note = "Abbreviations: CCO, continuous cardiac output; CO, cardiac output; CPO, cardiac power output; MAP, mean arterial pressure; mPAP, mean pulmonary artery pressure; ScvO₂, central venous oxygen saturation; SvO₂, mixed venous oxygen saturation. L denotes the prediction landmark. Sources were mapped to the same scientific constructs but were not assumed to be measurement-identical across databases."
    return headers, rows, [0.85, 1.65, 2.25, 1.25, 1.8, 1.8], set(), note


def predictor_summary(name: str, definition: str) -> tuple[str, str]:
    if name in {"landmark_h", "landmark_h_sq"}:
        return "At L", "Recorded value"
    if "ICU admission" in definition:
        return "ICU admission", "Recorded value"
    if "(L-4 h, L]" in definition:
        lookback = "(L-4 h, L]"
    else:
        lookback = "At L"
    summary_map = {
        "_mean_": "Mean", "_min_": "Minimum", "_max_": "Maximum", "_sd_": "Population SD",
        "_slope_": "Linear slope", "_delta_": "Last minus first", "_span_": "First-to-last span",
        "_current": "Most recent value", "recency": "Time since most recent value",
    }
    summary = "Indicator at landmark" if ("active_at_landmark" in name) else "Recorded value"
    if name.startswith("n_"):
        summary = "Measurement count"
    for key, val in summary_map.items():
        if key in name or name.startswith(key):
            summary = val
            break
    return lookback, summary


def make_table_s2() -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    d = pd.read_csv(PREDICTOR_FILE)
    rows: list[list[str]] = []
    for _, r in d.iterrows():
        block = r["block"]
        model = "A, B, C" if block == "shared/Model A" else ("B, C" if block == "Model B addition" else "C")
        predictor = str(r["predictor"])
        definition = str(r["definition_and_timing"])
        if predictor == "mpap_current":
            definition = "Most recent mPAP in (L-4 h, L]"
        elif predictor == "svo2_current":
            definition = "Most recent SvO₂ in (L-4 h, L]"
        definition = definition.replace("SvO2", "SvO₂")
        lookback, summary = predictor_summary(predictor, definition)
        if str(r["unit"]) == "category":
            handling = "Development mode"
        elif "active_at_landmark" in predictor:
            handling = "Development median if missing; no MIMIC missingness"
        else:
            handling = "Development median"
        rows.append([predictor, model, definition, str(r["unit"]), lookback, summary, handling])
    headers = ["Predictor", "Model", "Definition", "Unit", "Lookback", "Summary", "Missing-data handling"]
    note = "Numeric predictors were median-imputed and categorical predictors were most-frequent-category imputed using parameters estimated in the MIMIC-IV development cohort. Categorical predictors were one-hot encoded; unseen validation categories were ignored by the frozen encoder. SD denotes standard deviation."
    return headers, rows, [1.35, 0.65, 2.75, 0.65, 1.0, 1.1, 1.55], set(), note


def weight_summary_rows(weight_data: dict[str, pd.DataFrame]) -> list[list[str]]:
    result = []
    for name in ["Development", "Temporal", "eICU"]:
        d = weight_data[name]
        obs_col = "primary_outcome_observed_6h" if name != "eICU" else "outcome_observed_6h"
        observed = d[d[obs_col].astype(bool)]
        raw = observed["ipow_stabilized"].dropna().to_numpy()
        trunc = observed["ipow_stabilized_truncated"].dropna().to_numpy()
        q01, q25, q50, q75, q99 = np.quantile(raw, [0.01, 0.25, 0.50, 0.75, 0.99])
        ess = (trunc.sum() ** 2) / np.square(trunc).sum()
        label = {"Development": "MIMIC-IV development", "Temporal": "MIMIC-IV temporal", "eICU": "eICU external"}[name]
        result.append([
            label,
            f"{100 * observed.shape[0] / d.shape[0]:.1f}",
            f"{q50:.3f} ({q25:.3f}-{q75:.3f})",
            f"{q01:.3f}-{q99:.3f}",
            f"{raw.max():.3f}",
            f"{trunc.max():.3f}",
            f"{ess:,.1f}",
            f"{100 * ess / len(trunc):.1f}",
        ])
    return result


def make_table_s3(weight_data) -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    headers = ["Cohort", "Observation rate, %", "Weight, median (IQR)", "q01-q99", "Max before truncation", "Max after truncation", "ESS", "ESS / observed N, %"]
    note = "ESS, effective sample size; IQR, interquartile range; IPOW, inverse probability-of-observation weighting. q01 and q99 denote the 1st and 99th percentiles of the stabilised weight distribution. MIMIC-IV weights used the development q01/q99 limits for both development and temporal evaluation; eICU used its prespecified cross-fitted observation model and locked positivity remediation."
    return headers, weight_summary_rows(weight_data), [1.35, 1.0, 1.35, 0.9, 1.25, 1.25, 0.75, 1.2], set(range(1, 8)), note


def make_table_s4() -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    m = pd.read_csv(MAIN_METRICS_FILE)
    d = pd.read_csv(MAIN_DELTAS_FILE)
    labels = [("brier", "Brier score", 6), ("auroc", "AUROC", 4), ("average_precision", "Average precision", 4), ("log_loss", "Log loss", 6)]
    rows = []
    for metric, label, dec in labels:
        b = float(m[(m.model == "B") & (m.metric == metric)].estimate.iloc[0])
        c = float(m[(m.model == "C") & (m.metric == metric)].estimate.iloc[0])
        dr = d[(d.comparison == "C-B") & (d.metric == metric)].iloc[0]
        rows.append([label, f"{b:.{dec}f}", f"{c:.{dec}f}", f"{dr.estimate:+.{dec}f} ({dr.ci_low:+.{dec}f} to {dr.ci_high:+.{dec}f})"])
    headers = ["Metric", "Model B", "Model C", "C-B (95% CI)"]
    note = "CI, confidence interval. Negative differences favour Model C for Brier score and log loss; positive differences favour Model C for AUROC and average precision. Intervals are percentile 95% CIs from 1,000 paired patient-cluster bootstrap samples."
    return headers, rows, [2.0, 1.35, 1.35, 3.25], {1, 2, 3}, note


def make_table_s5(sens: list[dict[str, object]]) -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    rows = []
    for r in sens:
        rows.append([
            str(r["label"]), f"{int(r['landmarks']):,}", f"{int(r['events']):,}", f"{int(r['patients']):,}",
            fmt_delta(r["brier"], 6), fmt_delta(r["auroc"], 4), fmt_delta(r["average_precision"], 4), fmt_delta(r["log_loss"], 6),
        ])
    e = pd.read_csv(EICU_ABS_FILE)
    e = e[e.analysis == "sensitivity_unweighted"]
    vals = {(r.model, r.metric): float(r.estimate) for _, r in e.iterrows()}
    rows.append([
        "eICU unweighted sensitivity", "7,753", "469", "157",
        f"{vals[('B','brier')] - vals[('A','brier')]:+.6f} (—)",
        f"{vals[('B','auroc')] - vals[('A','auroc')]:+.4f} (—)",
        f"{vals[('B','average_precision')] - vals[('A','average_precision')]:+.4f} (—)",
        f"{vals[('B','log_loss')] - vals[('A','log_loss')]:+.6f} (—)",
    ])
    headers = ["Analysis", "Landmarks", "Event landmarks", "Event patients", "Delta Brier (95% CI)", "Delta AUROC (95% CI)", "Delta AP (95% CI)", "Delta log loss (95% CI)"]
    note = "AP, average precision; CI, confidence interval; CPO, cardiac power output. All differences are Model B minus Model A. Negative values favour Model B for Brier score and log loss; positive values favour Model B for AUROC and AP. A dash indicates that a paired interval was not available in the frozen output; no new post hoc bootstrap was performed. The primary analysis used the common mPAP/SvO₂ risk set and IPOW."
    return headers, rows, [1.85, 0.75, 0.85, 0.82, 1.35, 1.35, 1.25, 1.35], set(range(1, 8)), note


def make_table_s6() -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    d = pd.read_csv(HOSPITAL_FILE)
    d = d.sort_values("delta_brier_B_minus_A", ascending=False)
    rows = []
    for _, r in d.iterrows():
        patients = int(r.event_patients + r.nonevent_patients)
        rows.append([
            f"Hospital {int(r.hospitalid)}", f"{patients:,}", f"{int(r.event_patients):,}", f"{int(r.observed_landmarks):,}", f"{int(r.event_landmarks):,}",
            f"{r.brier_A:.5f}", f"{r.brier_B:.5f}", f"{r.delta_brier_B_minus_A:+.5f}",
        ])
    headers = ["Hospital", "Patients", "Event patients", "Landmarks", "Event landmarks", "Model A Brier", "Model B Brier", "Delta Brier"]
    note = "Hospitals are ordered from the most positive to the most negative Brier-score difference, matching main Fig. 4C. Delta Brier is Model B minus Model A; negative values favour Model B. Hospital-specific analyses were descriptive and were restricted to hospitals with at least five event patients; no hospital-specific P values or confidence intervals were calculated."
    return headers, rows, [1.35, 0.85, 0.95, 0.9, 1.0, 1.15, 1.15, 1.05], set(range(1, 8)), note


def make_table_s7() -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    d = pd.read_csv(ALERT_FILE)
    d = d[d.model == "B"].sort_values("threshold")
    rows = []
    for _, r in d.iterrows():
        label = f"{r.threshold:.3f}" + ("†" if abs(r.threshold - 0.05) < 1e-9 else "")
        rows.append([label, f"{100*r.event_pair_sensitivity:.1f}", f"{100*r.event_patient_sensitivity:.1f}", f"{100*r.event_free_landmark_alert_fraction:.1f}", f"{100*r.no_event_patient_any_alert_fraction:.1f}"])
    headers = ["Threshold", "Event-pair sensitivity, %", "Event-patient sensitivity, %", "Event-free landmark alerts, %", "Event-free patients with ≥1 alert, %"]
    note = "† Prespecified descriptive threshold used for the main alert-timing analysis. Event pairs denote independent confirmed/repeated low-CPO episodes. Event-free patients were patients without any event-positive landmark during temporal evaluation. Threshold analyses were descriptive; no threshold was selected as optimal."
    return headers, rows, [1.0, 1.8, 1.8, 2.0, 2.15], set(range(1, 5)), note


def make_table_s8() -> tuple[list[str], list[list[str]], list[float], set[int], str]:
    rows = [
        ["Primary confirmed/repeated CPO <0.60 W within 6 h", "Two values <0.60 W separated by 0.5-3 h", "4,730", "202", "62", "Yes, Models A/B/C", "Primary"],
        ["Confirmed/repeated CPO <0.60 W within 12 h", "Same confirmation rule within 12 h", "4,583", "319", "65", "No", "Counts only"],
        ["Confirmed/repeated CPO <0.50 W within 6 h", "Two values <0.50 W separated by 0.5-3 h", "4,730", "38", "16", "No", "Event count insufficient"],
        ["Confirmed/repeated CPO <0.50 W within 12 h", "Same confirmation rule within 12 h", "4,583", "66", "19", "No", "Counts only"],
        ["Low CPO with concurrent SvO₂ <60%", "No prespecified temporal matching rule", "—", "—", "—", "No", "Not operationalised"],
    ]
    headers = ["Outcome definition", "Confirmation rule", "Landmarks", "Event landmarks", "Event patients", "Model fitted", "Decision"]
    note = "Landmark counts in this table refer to the temporal legal-landmark common risk set before restriction to outcome-observed landmarks. The primary predictive analysis included 3,416 landmarks with observed 6-h outcomes; 202 met the confirmed/repeated low-CPO event definition. Only the 6-h confirmed/repeated CPO <0.60-W outcome was modelled. The 0.50-W and 12-h definitions had been declared as sensitivity counts only. The joint low-CPO/SvO₂ definition was not operationalised because a temporal matching rule had not been prespecified; choosing one after outcome inspection would have introduced post hoc analytical flexibility."
    return headers, rows, [2.25, 2.0, 0.75, 0.9, 0.85, 1.15, 1.3], {2, 3, 4}, note


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size in [("Title", 20), ("Subtitle", 11), ("Heading 1", 13), ("Heading 2", 10.8)]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(18)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(11)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    title_p_pr = doc.styles["Title"]._element.get_or_add_pPr()
    for border in title_p_pr.findall(qn("w:pBdr")):
        title_p_pr.remove(border)


def remove_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for border in p_pr.findall(qn("w:pBdr")):
        p_pr.remove(border)


def add_methods(doc: Document) -> None:
    doc.add_heading("Supplementary Methods", level=1)

    methods = [
        ("S1 Database mapping and haemodynamic extraction", [
            "MIMIC-IV version 3.1 and eICU Collaborative Research Database version 2.0 were mapped to prespecified haemodynamic constructs before clinical model evaluation. In MIMIC-IV, cardiac output was obtained from charted continuous cardiac-output item 224842. Mean arterial pressure was taken from the prespecified hierarchy of invasive or non-invasive MAP items 220052, 225312, and 220181. Cardiac power output was calculated as cardiac output multiplied by MAP and divided by 451. Mean pulmonary artery pressure used item 220061. Mixed venous oxygen saturation used item 223772; central venous oxygen saturation item 226541 was not substituted.",
            "In eICU, the transported cardiac-output source was nurseCharting with exact label and name CO|CO. Each CO value was paired using a frozen hierarchy: exact-offset nurse-charted invasive MAP, otherwise the most recent vitalPeriodic systemicMean within five minutes, and otherwise exact-offset vitalAperiodic noninvasiveMean. Values after the CO offset were never used. eICU mPAP and SvO₂ used exact nurseCharting labels PA|PA Mean and SVO2|SVO2. Because eICU does not distinguish continuous thermodilution, intermittent thermodilution, and Fick cardiac-output measurements, external validation was interpreted as transport under a related measurement process rather than measurement-identical validation.",
            "For both databases, implausible values were removed using prespecified physiological ranges, identical duplicates were collapsed, and conflicting valid values recorded for the same stay, source, and time were excluded. Haemodynamic observations were required to occur within the ICU stay and to satisfy the time rule specified for the corresponding analysis. Table S1 gives the cross-database mapping."
        ]),
        ("S2 Temporal split reconstruction", [
            "MIMIC-IV shifts calendar dates within patient. For each hospital stay, we reconstructed a possible source-year interval from anchor_year, anchor_year_group, and the difference between the recorded admission year and anchor_year. Patients whose latest possible source year was 2019 or earlier formed the development cohort; patients whose earliest possible source year was 2020 or later formed the temporal-evaluation cohort. Patients whose possible intervals crossed the 2019/2020 boundary were excluded from development and primary temporal evaluation.",
            "All ICU stays from a patient were retained in the same period. The split was determined before fitting clinical outcome models. The development period supplied imputation, encoding, model-fitting, and calibration parameters; the temporal period was used only for locked evaluation."
        ]),
        ("S3 Landmark and information clock implementation", [
            "Hourly landmarks were generated from ICU hour 6 through hour 66. A landmark required a most recent CPO of at least 0.60 W within the preceding hour, at least three CPO measurements in the preceding four hours, a CPO observation span of at least two hours, and no previously confirmed/repeated low-CPO episode. Major mechanical circulatory support active at the landmark was excluded, whereas intra-aortic balloon pump status was retained as a predictor. Major support initiated after the landmark and before outcome ascertainment triggered censoring.",
            "The primary MIMIC analysis used charttime as the closest available proxy for physiological measurement time. A strict information-clock sensitivity required the derived CPO record to have been stored by the landmark, using the later storetime of its cardiac-output and MAP components. mPAP and SvO₂ records also required storetime no later than the landmark. All haemodynamic features and observation weights were rebuilt under this restriction without changing the estimator, hyperparameters, temporal split, or outcome definition.",
            "In eICU, nurse-charted observations were available only after nursingChartEntryOffset. Monitor-interface values without an entry offset were treated as available at their observation offset. Dynamic histories used the interval (L-4 h, L], and no future observation, treatment, or outcome information was used as a predictor."
        ]),
        ("S4 Outcome observation and weighting", [
            "The primary outcome was a confirmed/repeated low-CPO episode within six hours after the landmark, defined as two CPO values below 0.60 W separated by 0.5-3 hours. A normal value between the two low values was allowed. An event-free outcome required at least four future CPO measurements, a final measurement at least five hours after the landmark, and no interval greater than two hours from the current value through the future sequence. Landmarks meeting neither the positive nor dense-negative definition had an unobserved outcome. For eICU, only CPO measurements at or before unitDischargeOffset were eligible for outcome classification. Loss of subsequent CPO measurements, including after ICU exit or discontinuation of haemodynamic monitoring, could therefore yield an outcome-unobserved landmark. Death and pulmonary artery catheter removal were not coded as separate censoring events; major mechanical circulatory support initiated before outcome confirmation was the prespecified explicit censoring event.",
            "Inverse probability-of-observation weighting was used to mitigate incomplete six-hour outcome ascertainment. In MIMIC-IV, an L2-regularised logistic model predicted outcome observation using prespecified landmark-time predictors. Five-fold stratified patient-grouped cross-validation supplied development out-of-fold probabilities; a model fitted to all development patients was applied unchanged to the temporal period. Stabilised weights used the development observation fraction divided by predicted observation probability and were truncated at the first and 99th percentiles of the development distribution.",
            "The eICU observation model used five-fold stratified cross-fitting grouped by hospital. Two observations fell below the prespecified positivity threshold of 0.05; before clinical predictions were generated, probabilities were floored at 0.05 and weights were truncated at the external first and 99th percentiles. Weight distributions, truncation limits, and effective sample size are reported in Table S3 and Fig. S1."
        ]),
        ("S5 Model implementation", [
            "All outcome models used histogram gradient boosting with fixed hyperparameters: learning rate 0.05, 200 iterations, at most 15 leaf nodes, maximum depth 3, at least 100 observations per leaf, L2 regularisation 1.0, no early stopping, and random seed 20260908. Numeric predictors were median-imputed, categorical predictors were mode-imputed and one-hot encoded, and all preprocessing parameters were estimated in development data. No class weighting, over-sampling, under-sampling, synthetic observations, outcome-driven feature selection, or temporal-period recalibration was used.",
            "Model A contained shared clinical and observation context plus the latest CPO. Model B added the four-hour CPO mean, minimum, maximum, standard deviation, linear slope, first-to-last change, and observed time span. Model C added the current value and corresponding four-hour summaries for mPAP and SvO₂. Models B and C were compared on the same common risk set. Development-only five-fold patient-grouped out-of-fold predictions defined a logistic calibration map for each model, which was then applied unchanged to temporal and external predictions. Table S2 provides the complete predictor dictionary."
        ]),
        ("S6 Bootstrap and sensitivity analyses", [
            "The primary metric was the IPOW-weighted temporal Brier score. Secondary metrics were weighted AUROC, average precision, log loss, calibration intercept, and calibration slope. A and B, and B and C, were compared on identical landmarks with identical weights. Percentile 95% confidence intervals used 1,000 paired patient-cluster bootstrap samples; patient multiplicity modified landmark weights, and individual landmarks were not resampled independently.",
            "Prespecified and locked sensitivity analyses excluded landmarks with a prior CPO below 0.60 W, restricted the latest CPO to at least 0.70 W, applied the strict EHR information clock, expanded A-versus-B evaluation to the broader CPO-eligible risk set, and repeated evaluation with unit weights. Frozen Models A and B were applied to eICU without refitting or recalibration; external intervals used 1,000 paired two-level bootstrap samples, first resampling hospitals and then patients within hospitals.",
            "Threshold analyses were descriptive and used the locked Model B temporal predictions. Event-pair and event-patient sensitivity, event-free landmark alert burden, and the proportion of event-free patients receiving at least one alert were reported at thresholds 0.020, 0.030, 0.050, 0.075, 0.100, and 0.150. No threshold was selected as optimal. Alternative 0.50-W and 12-hour endpoints were retained as prespecified counts only and were not used to fit additional models."
        ]),
    ]
    for heading, paragraphs in methods:
        doc.add_heading(heading, level=2)
        for text in paragraphs:
            para = doc.add_paragraph(text)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def build_doc(figures: list[Path], weight_data, sens) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    configure_styles(doc)
    add_page_number(sec.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title)
    r = title.add_run("Recent cardiac power history improves six-hour prediction of low-CPO episodes during pulmonary artery catheter monitoring")
    set_font(r, size=17, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Supplementary Information")
    set_font(r, size=12, bold=True)
    intro = doc.add_paragraph(
        "This Supplementary Information provides the database mapping, complete predictor definitions, observation-weight diagnostics, prespecified sensitivity analyses, hospital-specific external-validation results, alert-threshold analyses, and alternative outcome counts supporting the development, temporal evaluation, and multicentre external validation of the CPO prediction models."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_page_break()
    add_methods(doc)

    add_landscape_section(doc)
    table_specs = [
        ("Table S1", "Haemodynamic variable definitions and database mapping", make_table_s1()),
        ("Table S2", "Complete predictor definitions and time windows", make_table_s2()),
        ("Table S3", "Outcome observation model and weighting diagnostics", make_table_s3(weight_data)),
        ("Table S4", "Incremental value of mPAP and SvO₂ history beyond CPO history", make_table_s4()),
        ("Table S5", "Prespecified sensitivity analyses of Model B versus Model A", make_table_s5(sens)),
        ("Table S6", "Hospital-specific external validation results", make_table_s6()),
        ("Table S7", "Alert-threshold analysis", make_table_s7()),
        ("Table S8", "Alternative outcome definitions and event counts", make_table_s8()),
    ]
    for idx, (number, title_text, spec) in enumerate(table_specs):
        if idx > 0:
            doc.add_page_break()
        headers, rows, widths, numeric_cols, note = spec
        add_caption(doc, number, title_text)
        bold_rows = {0} if number == "Table S5" else set()
        size = 6.9 if number in {"Table S1", "Table S2", "Table S5"} else 7.4
        add_nature_table(doc, headers, rows, widths, numeric_cols, bold_rows, size)
        add_note(doc, note)

    figure_titles = [
        ("Figure S1", "Outcome observation and weight diagnostics"),
        ("Figure S2", "Robustness of the incremental value of recent CPO history"),
        ("Figure S3", "Incremental value of additional mPAP and SvO₂ history beyond recent CPO history"),
    ]
    figure_legends = [
        "a-c, Distributions of cross-fitted predicted probabilities that the 6-h outcome would be observed in the MIMIC-IV development, MIMIC-IV temporal, and eICU external cohorts. d-f, Raw and q01/q99-truncated stabilised observation weights among outcome-observed landmarks. Dashed vertical lines mark the raw q01 and q99 limits. ESS is calculated from the truncated weights and is shown as the absolute effective sample size and percentage of observed landmarks.",
        "a, Paired Brier-score differences and b, paired AUROC differences for Model B versus Model A across the primary temporal analysis and prespecified or locked sensitivity analyses. Points show differences and horizontal lines show percentile 95% confidence intervals from 1,000 paired patient-cluster bootstrap samples. Larger points identify the primary temporal analysis. The open point for latest CPO ≥0.70 W is the frozen point estimate; no paired interval was available and no post hoc bootstrap was performed.",
        "Model C minus Model B paired differences for Brier score, AUROC, average precision, and log loss in the locked temporal cohort. Each mini-panel uses its own numerical scale; points show differences and horizontal lines show percentile 95% confidence intervals from 1,000 paired patient-cluster bootstrap samples. The vertical line denotes no difference. Negative differences favour Model C for Brier score and log loss, whereas positive differences favour Model C for AUROC and average precision.",
    ]
    for i, ((number, title_text), fig, legend) in enumerate(zip(figure_titles, figures, figure_legends)):
        doc.add_page_break()
        add_caption(doc, number, title_text)
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width = Inches(10.1 if i == 0 else (9.25 if i == 1 else 9.0))
        para.add_run().add_picture(str(fig), width=width)
        add_note(doc, legend + " Source data are provided in the reproducibility package.")

    props = doc.core_properties
    props.title = "Supplementary Information for the CPO history prediction study"
    props.subject = "Supplementary methods, tables, and figures"
    props.author = ""
    props.keywords = "cardiac power output; temporal validation; external validation; supplementary information"
    OUTDIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTDOC)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["all", "figures", "docx"], default="all")
    args = parser.parse_args()
    weight_data = prepare_weight_data()
    sens = sensitivity_rows()
    figures = [
        FIGDIR / "figure_S1_observation_weight_diagnostics.png",
        FIGDIR / "figure_S2_robustness_forest.png",
        FIGDIR / "figure_S3_model_c_incremental_value.png",
    ]
    if args.mode in {"all", "figures"}:
        figures = [fig_s1(weight_data), fig_s2(sens), fig_s3()]
        for f in figures:
            print(f)
    if args.mode in {"all", "docx"}:
        missing = [f for f in figures if not f.exists()]
        if missing:
            raise FileNotFoundError(f"Missing generated figures: {missing}")
        build_doc(figures, weight_data, sens)
        print(OUTDOC)


if __name__ == "__main__":
    main()
